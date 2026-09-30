from __future__ import annotations

from datetime import datetime
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from category_registry import normalize_category_name, normalize_feeling
from config import (
    CATEGORY_CONFIG_COLUMNS,
    CATEGORY_CONFIG_PATH,
    CATEGORIES,
    CORE_CATEGORY_SET,
    DAILY_SUMMARY_COLUMNS,
    DAILY_SUMMARY_FILE_PREFIX,
    DATA_BACKUP_DIR,
    DATA_DIR,
    DRAFT_DIR,
    DEFAULT_CATEGORY_CONFIG_ROWS,
    LEGACY_TIME_LOG_PATH,
    LLM_PROMPT_FILE_PREFIX,
    OUTPUT_DIR,
    OUTPUT_CHART_DIR,
    OUTPUT_BACKUP_ENABLED,
    ROOT_LEGACY_TIME_LOG_PATH,
    SCORE_DETAIL_COLUMNS,
    TIME_LOG_COLUMNS,
    TIME_LOG_EXPORT_COLUMNS_ZH,
    TIME_LOG_FILE_PREFIX,
    TODAY_REPORT_FILE_PREFIX,
)


LOG_IMPORT_COLUMNS_ZH_TO_EN = {
    value: key for key, value in TIME_LOG_EXPORT_COLUMNS_ZH.items()
}

MONTH_DIR_PATTERN = re.compile(r"^\d{4}\.\d{2}$")
DAY_DIR_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _month_tag_from_date(date_str: str) -> str:
    date_obj = datetime.strptime(date_str, "%Y-%m-%d")
    return date_obj.strftime("%Y.%m")


def get_data_month_dir(date_str: str) -> Path:
    return DATA_DIR / _month_tag_from_date(date_str)


def get_data_day_dir(date_str: str) -> Path:
    return get_data_month_dir(date_str) / date_str


def get_daily_raw_log_path(date_str: str) -> Path:
    return get_data_day_dir(date_str) / f"{TIME_LOG_FILE_PREFIX}_{date_str}.xlsx"


def get_output_month_dir(date_str: str) -> Path:
    return OUTPUT_DIR / _month_tag_from_date(date_str)


def get_output_day_dir(date_str: str) -> Path:
    return get_output_month_dir(date_str) / date_str


def get_output_month_chart_dir(date_str: str) -> Path:
    return get_output_month_dir(date_str) / "charts"


def get_output_chart_dir() -> Path:
    return OUTPUT_CHART_DIR


def get_daily_summary_output_path(date_str: str) -> Path:
    return get_output_day_dir(date_str) / f"{DAILY_SUMMARY_FILE_PREFIX}_{date_str}.xlsx"


def get_today_report_output_path(date_str: str) -> Path:
    return get_output_day_dir(date_str) / f"{TODAY_REPORT_FILE_PREFIX}_{date_str}.txt"


def get_llm_prompt_output_path(date_str: str) -> Path:
    return get_output_day_dir(date_str) / f"{LLM_PROMPT_FILE_PREFIX}_{date_str}.txt"


def get_daily_output_paths(date_str: str) -> list[Path]:
    return [
        get_daily_summary_output_path(date_str),
        get_today_report_output_path(date_str),
        get_llm_prompt_output_path(date_str),
    ]


def get_backup_version_dir(now_dt: datetime | None = None) -> Path:
    now = now_dt or datetime.now()
    version_tag = f"v{now.year}.{now.month}.{now.day}"
    return DATA_BACKUP_DIR / version_tag


def backup_file_if_exists(path: Path, now_dt: datetime | None = None) -> Path | None:
    if not path.exists():
        return None

    backup_dir = get_backup_version_dir(now_dt)
    backup_dir.mkdir(parents=True, exist_ok=True)

    target = backup_dir / path.name
    if target.exists():
        ts = (now_dt or datetime.now()).strftime("%H%M%S")
        target = backup_dir / f"{path.stem}_{ts}{path.suffix}"
        index = 1
        while target.exists():
            target = backup_dir / f"{path.stem}_{ts}_{index}{path.suffix}"
            index += 1

    shutil.copy2(path, target)
    return target


class DataReadError(RuntimeError):
    """A source workbook could not be read without risking incomplete statistics."""


def _safe_file_label(path: Path) -> str:
    """Show a project relative path without exposing an external home directory."""
    try:
        return path.resolve().relative_to(DATA_DIR.resolve().parent).as_posix()
    except ValueError:
        return path.name


def _temporary_path(target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{target.stem}-", suffix=target.suffix,
        dir=target.parent, delete=False,
    ) as stream:
        return Path(stream.name)


def _remove_temporary(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass  # Preserve the original write failure.


def _backup_daily_source(path: Path, date_str: str) -> None:
    """Keep one previous successful daily workbook, replacing the backup atomically."""
    backup = DATA_BACKUP_DIR / _month_tag_from_date(date_str) / date_str / path.name
    temporary = _temporary_path(backup)
    try:
        shutil.copy2(path, temporary)
        # The source was read before writing; verify the copied file is still readable.
        with pd.ExcelFile(temporary):
            pass
        os.replace(temporary, backup)
    finally:
        _remove_temporary(temporary)


def _read_csv_with_fallback(path: Path) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(path)


def _default_category_config_df() -> pd.DataFrame:
    return pd.DataFrame(DEFAULT_CATEGORY_CONFIG_ROWS, columns=CATEGORY_CONFIG_COLUMNS)


def _normalize_mode(value: Any, default_value: str) -> str:
    allowed = {"maximize", "target", "minimize", "range", "record_only"}
    text = str(value).strip().lower() if value is not None else default_value
    return text if text in allowed else default_value


def _normalize_numeric(value: Any, default_value: float) -> float:
    parsed = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(parsed):
        return float(default_value)
    return float(parsed)


def _normalize_category_config_df(raw_df: pd.DataFrame) -> pd.DataFrame:
    default_df = _default_category_config_df()
    default_map = {row["category"]: row.to_dict() for _, row in default_df.iterrows()}

    if raw_df.empty or "category" not in raw_df.columns:
        return default_df

    input_df = raw_df.copy()
    input_df["category"] = input_df["category"].map(normalize_category_name)
    input_df = input_df[input_df["category"].isin(CATEGORIES)]
    input_df = input_df.drop_duplicates(subset=["category"], keep="last")
    input_map = {row["category"]: row.to_dict() for _, row in input_df.iterrows()}

    rows: list[dict[str, Any]] = []
    for category in CATEGORIES:
        default_row = default_map[category]
        input_row = input_map.get(category, {})
        rows.append(
            {
                "category": category,
                "mode": _normalize_mode(input_row.get("mode"), default_row["mode"]),
                "target_min": _normalize_numeric(
                    input_row.get("target_min"), default_row["target_min"]
                ),
                "lower_min": _normalize_numeric(
                    input_row.get("lower_min"), default_row["lower_min"]
                ),
                "upper_min": _normalize_numeric(
                    input_row.get("upper_min"), default_row["upper_min"]
                ),
                "weight": _normalize_numeric(input_row.get("weight"), default_row["weight"]),
                "core": 1 if category in CORE_CATEGORY_SET else 0,
            }
        )
    return pd.DataFrame(rows, columns=CATEGORY_CONFIG_COLUMNS)


def ensure_category_config_file() -> None:
    if not CATEGORY_CONFIG_PATH.exists():
        temporary = _temporary_path(CATEGORY_CONFIG_PATH)
        try:
            _default_category_config_df().to_csv(temporary, index=False, encoding="utf-8-sig")
            if list(_read_csv_with_fallback(temporary).columns) != CATEGORY_CONFIG_COLUMNS:
                raise ValueError("default category configuration schema mismatch")
            os.replace(temporary, CATEGORY_CONFIG_PATH)
        finally:
            _remove_temporary(temporary)
        return

    raw_df = _read_csv_with_fallback(CATEGORY_CONFIG_PATH)
    if "category" not in raw_df.columns:
        raise ValueError(f"{_safe_file_label(CATEGORY_CONFIG_PATH)}: missing category column")
    # Normalize in memory on load. Never rewrite a user's configuration on startup.


def load_category_config() -> pd.DataFrame:
    ensure_category_config_file()
    raw_df = _read_csv_with_fallback(CATEGORY_CONFIG_PATH)
    return _normalize_category_config_df(raw_df)


def _normalize_log_df(df: pd.DataFrame, fallback_date: str | None = None) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=TIME_LOG_COLUMNS)

    out = df.copy()
    out = out.rename(columns=LOG_IMPORT_COLUMNS_ZH_TO_EN)
    for col in TIME_LOG_COLUMNS:
        if col not in out.columns:
            out[col] = ""

    if fallback_date:
        out["date"] = out["date"].replace("", fallback_date).fillna(fallback_date)

    out["date"] = pd.to_datetime(out["date"], errors="coerce").dt.strftime("%Y-%m-%d")
    if fallback_date:
        out["date"] = out["date"].fillna(fallback_date)

    out = out.dropna(subset=["date"])
    out = out[out["date"].astype(str).str.len() > 0]

    out["category"] = out["category"].map(normalize_category_name)
    out["duration_min"] = pd.to_numeric(out["duration_min"], errors="coerce").fillna(0).astype(int)
    out["duration_min"] = out["duration_min"].clip(lower=0)
    out["feeling"] = out["feeling"].apply(normalize_feeling)
    out["note"] = out["note"].fillna("").astype(str)
    out["session_id"] = out["session_id"].fillna("").astype(str)
    out["created_at"] = out["created_at"].fillna("").astype(str)
    return out[TIME_LOG_COLUMNS]


def _fit_worksheet_style(worksheet, left_headers: set[str] | None = None) -> None:
    if left_headers is None:
        left_headers = set()

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    worksheet.freeze_panes = "A2"

    headers: list[str] = []
    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = center
        headers.append(str(cell.value or ""))

    max_col = worksheet.max_column
    max_row = worksheet.max_row
    for col_idx in range(1, max_col + 1):
        header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
        col_letter = get_column_letter(col_idx)
        max_len = len(header)
        for row_idx in range(2, max_row + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = left if header in left_headers else center
            txt = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(txt))
        worksheet.column_dimensions[col_letter].width = min(max(10, max_len + 2), 50)


def _write_daily_raw_log_file(
    date_str: str,
    day_df: pd.DataFrame,
    backup_before_overwrite: bool = True,
) -> Path:
    target_path = get_daily_raw_log_path(date_str)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "时间记录"

    export_df = day_df.rename(columns=TIME_LOG_EXPORT_COLUMNS_ZH)[
        list(TIME_LOG_EXPORT_COLUMNS_ZH.values())
    ]
    worksheet.append(export_df.columns.tolist())
    for row in export_df.itertuples(index=False, name=None):
        worksheet.append(list(row))

    _fit_worksheet_style(worksheet, left_headers={"备注"})
    temporary = _temporary_path(target_path)
    try:
        workbook.save(temporary)
        with pd.ExcelFile(temporary) as check:
            if check.sheet_names != ["时间记录"]:
                raise ValueError("daily workbook sheet mismatch")
            written = pd.read_excel(check, sheet_name="时间记录")
        if list(written.columns) != list(TIME_LOG_EXPORT_COLUMNS_ZH.values()) or len(written) != len(day_df):
            raise ValueError("daily workbook validation failed")
        if backup_before_overwrite and target_path.exists():
            _backup_daily_source(target_path, date_str)
        os.replace(temporary, target_path)
    finally:
        _remove_temporary(temporary)
    return target_path


def save_daily_segments(
    date_str: str,
    segments: list[dict[str, Any]],
    session_id: str,
) -> tuple[pd.DataFrame, Path]:
    rows = []
    for item in segments:
        created_at = str(item.get("created_at", "")).strip() or datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S.%f"
        )
        rows.append(
            {
                "date": date_str,
                "category": normalize_category_name(item["category"]),
                "duration_min": int(item["duration_min"]),
                "feeling": normalize_feeling(item["feeling"]),
                "note": str(item.get("note", "")),
                "session_id": str(item.get("session_id", "")).strip() or session_id,
                "created_at": created_at,
            }
        )

    new_df = pd.DataFrame(rows, columns=TIME_LOG_COLUMNS)
    new_df = _normalize_log_df(new_df, fallback_date=date_str)

    target_path = get_daily_raw_log_path(date_str)
    if target_path.exists():
        existing_raw = pd.read_excel(target_path)
        existing_df = _normalize_log_df(existing_raw, fallback_date=date_str)
        day_df = pd.concat([existing_df, new_df], ignore_index=True)
        day_df = day_df.drop_duplicates(subset=TIME_LOG_COLUMNS, keep="first").reset_index(drop=True)
    else:
        day_df = new_df

    _write_daily_raw_log_file(
        date_str=date_str,
        day_df=day_df,
        backup_before_overwrite=True,
    )
    return day_df, target_path


def backup_daily_output_files(date_str: str) -> Path | None:
    if not OUTPUT_BACKUP_ENABLED:
        return None

    day_dir = get_output_day_dir(date_str)
    if not day_dir.exists():
        return None

    files = get_daily_output_paths(date_str)
    existing_files = [path for path in files if path.exists()]
    if not existing_files:
        return None

    index = 1
    while (day_dir / f"{date_str}_backup_{index}").exists():
        index += 1

    backup_dir = day_dir / f"{date_str}_backup_{index}"
    backup_dir.mkdir(parents=True, exist_ok=False)
    for path in existing_files:
        shutil.copy2(path, backup_dir / path.name)
    return backup_dir


def _iter_daily_log_files() -> list[Path]:
    files: list[Path] = []
    for month_dir in DATA_DIR.iterdir():
        if not month_dir.is_dir() or not MONTH_DIR_PATTERN.match(month_dir.name):
            continue
        for day_dir in month_dir.iterdir():
            if not day_dir.is_dir() or not DAY_DIR_PATTERN.match(day_dir.name):
                continue
            file_path = day_dir / f"{TIME_LOG_FILE_PREFIX}_{day_dir.name}.xlsx"
            if file_path.exists():
                files.append(file_path)
    return sorted(files)


def _read_daily_log_file(path: Path) -> pd.DataFrame:
    fallback_date = path.stem.split("_")[-1] if "_" in path.stem else None
    raw_df = pd.read_excel(path)
    column_names = set(raw_df.columns)
    if not {"类别", "时长（分钟）"}.issubset(column_names) and not {"category", "duration_min"}.issubset(column_names):
        raise ValueError("daily workbook is missing required columns")
    return _normalize_log_df(raw_df, fallback_date=fallback_date)


def _iter_legacy_monthly_log_files() -> list[Path]:
    files = [
        path for path in DATA_DIR.glob("*_时间记录.xlsx")
        if path.is_file()
    ]
    return sorted(files)


def _load_legacy_monthly_logs_df() -> pd.DataFrame:
    chunks: list[pd.DataFrame] = []
    for monthly_path in _iter_legacy_monthly_log_files():
        try:
            all_sheets = pd.read_excel(monthly_path, sheet_name=None)
        except Exception as exc:
            raise DataReadError(
                f"{_safe_file_label(monthly_path)}: {type(exc).__name__}: unable to read legacy workbook"
            ) from None
        for sheet_name, sheet_df in all_sheets.items():
            fallback_date = str(sheet_name) if DAY_DIR_PATTERN.match(str(sheet_name)) else None
            norm_df = _normalize_log_df(sheet_df, fallback_date=fallback_date)
            if not norm_df.empty:
                chunks.append(norm_df)

    if not chunks:
        return pd.DataFrame(columns=TIME_LOG_COLUMNS)
    return pd.concat(chunks, ignore_index=True)


def _load_legacy_time_log_df() -> pd.DataFrame:
    for path in (LEGACY_TIME_LOG_PATH, ROOT_LEGACY_TIME_LOG_PATH):
        if not path.exists():
            continue
        try:
            raw_df = pd.read_excel(path)
            norm_df = _normalize_log_df(raw_df)
            if not norm_df.empty:
                return norm_df
        except Exception as exc:
            raise DataReadError(
                f"{_safe_file_label(path)}: {type(exc).__name__}: unable to read legacy workbook"
            ) from None
    return pd.DataFrame(columns=TIME_LOG_COLUMNS)


def _migrate_legacy_logs_if_needed() -> None:
    if _iter_daily_log_files():
        return

    legacy_monthly = _iter_legacy_monthly_log_files()
    if legacy_monthly:
        for monthly_path in legacy_monthly:
            try:
                all_sheets = pd.read_excel(monthly_path, sheet_name=None)
            except Exception as exc:
                raise DataReadError(
                    f"{_safe_file_label(monthly_path)}: {type(exc).__name__}: unable to migrate legacy workbook"
                ) from None
            for sheet_name, sheet_df in all_sheets.items():
                fallback_date = str(sheet_name) if DAY_DIR_PATTERN.match(str(sheet_name)) else None
                norm_df = _normalize_log_df(sheet_df, fallback_date=fallback_date)
                if norm_df.empty:
                    continue
                date_str = str(norm_df.iloc[0]["date"])
                _write_daily_raw_log_file(
                    date_str=date_str,
                    day_df=norm_df,
                    backup_before_overwrite=False,
                )
        return

    legacy_df = _load_legacy_time_log_df()
    if legacy_df.empty:
        return
    for date_str, day_df in legacy_df.groupby("date", sort=True):
        _write_daily_raw_log_file(
            date_str=str(date_str),
            day_df=day_df[TIME_LOG_COLUMNS],
            backup_before_overwrite=False,
        )


def ensure_project_files() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DRAFT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_CHART_DIR.mkdir(parents=True, exist_ok=True)
    DATA_BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    ensure_category_config_file()
    _migrate_legacy_logs_if_needed()


def load_all_time_logs() -> pd.DataFrame:
    daily_files = _iter_daily_log_files()
    chunks: list[pd.DataFrame] = []
    date_covered: set[str] = set()

    for file_path in daily_files:
        try:
            df = _read_daily_log_file(file_path)
            if not df.empty:
                chunks.append(df)
                date_covered.update(df["date"].astype(str).tolist())
        except Exception as exc:
            raise DataReadError(
                f"{_safe_file_label(file_path)}: {type(exc).__name__}: unable to read daily workbook"
            ) from None

    legacy_monthly_df = _load_legacy_monthly_logs_df()
    if not legacy_monthly_df.empty:
        if date_covered:
            legacy_monthly_df = legacy_monthly_df[~legacy_monthly_df["date"].isin(date_covered)]
        if not legacy_monthly_df.empty:
            chunks.append(legacy_monthly_df)
            date_covered.update(legacy_monthly_df["date"].astype(str).tolist())

    legacy_single_df = _load_legacy_time_log_df()
    if not legacy_single_df.empty:
        if date_covered:
            legacy_single_df = legacy_single_df[~legacy_single_df["date"].isin(date_covered)]
        if not legacy_single_df.empty:
            chunks.append(legacy_single_df)

    if not chunks:
        return pd.DataFrame(columns=TIME_LOG_COLUMNS)

    out = pd.concat(chunks, ignore_index=True)
    out = out.sort_values(["date", "created_at"], kind="stable").reset_index(drop=True)
    return out[TIME_LOG_COLUMNS]


def _build_daily_overview_df(target_date: str, daily_summary_df: pd.DataFrame) -> pd.DataFrame:
    today_summary = daily_summary_df[daily_summary_df["date"] == target_date]
    if today_summary.empty:
        row = {
            "日期": target_date,
            "今日总记录分钟数": 0,
            "今日总记录小时数": 0.0,
            "今日时间段数量": 0,
            "今日总体效率分（0~100）": 0.0,
            "今日核心成长分": 0.0,
        }
    else:
        summary_row = today_summary.iloc[0]
        total_min = float(summary_row["total_recorded_min"])
        row = {
            "日期": target_date,
            "今日总记录分钟数": int(round(total_min)),
            "今日总记录小时数": round(total_min / 60.0, 2),
            "今日时间段数量": int(summary_row["segment_count"]),
            "今日总体效率分（0~100）": round(float(summary_row["total_score"]), 2),
            "今日核心成长分": round(float(summary_row["core_growth_score"]), 2),
        }
    return pd.DataFrame([row])


def _build_daily_detail_df(
    day_detail_df: pd.DataFrame,
    category_config_df: pd.DataFrame,
) -> pd.DataFrame:
    base_df = category_config_df[["category", "core"]].copy()
    metric_cols = [
        "category",
        "total_recorded_min",
        "segment_count",
        "total_effective_min",
        "category_score",
    ]
    metrics_df = day_detail_df[metric_cols].copy() if not day_detail_df.empty else pd.DataFrame(columns=metric_cols)
    detail_df = base_df.merge(metrics_df, on="category", how="left")

    detail_df["total_recorded_min"] = detail_df["total_recorded_min"].fillna(0.0)
    detail_df["segment_count"] = detail_df["segment_count"].fillna(0).astype(int)
    detail_df["total_effective_min"] = detail_df["total_effective_min"].fillna(0.0)
    detail_df["category_score"] = detail_df["category_score"].fillna(0.0)

    detail_df["今日总记录分钟数"] = detail_df["total_recorded_min"].round().astype(int)
    detail_df["今日总记录小时数"] = (detail_df["total_recorded_min"] / 60.0).round(2)
    detail_df["今日时间段数量"] = detail_df["segment_count"].astype(int)
    detail_df["今日有效分钟数"] = detail_df["total_effective_min"].round(2)
    detail_df["今日类别得分"] = detail_df["category_score"].round(2)
    detail_df["是否核心"] = detail_df["core"].apply(lambda value: "是" if int(value) == 1 else "否")

    detail_df = detail_df[
        [
            "category",
            "今日总记录分钟数",
            "今日总记录小时数",
            "今日时间段数量",
            "今日有效分钟数",
            "今日类别得分",
            "是否核心",
        ]
    ].rename(columns={"category": "类别"})

    total_row = pd.DataFrame(
        [
            {
                "类别": "总计",
                "今日总记录分钟数": int(detail_df["今日总记录分钟数"].sum()),
                "今日总记录小时数": round(float(detail_df["今日总记录小时数"].sum()), 2),
                "今日时间段数量": int(detail_df["今日时间段数量"].sum()),
                "今日有效分钟数": round(float(detail_df["今日有效分钟数"].sum()), 2),
                "今日类别得分": round(float(detail_df["今日类别得分"].sum()), 2),
                "是否核心": "-",
            }
        ]
    )
    return pd.concat([detail_df, total_row], ignore_index=True)


def _verify_saved_daily_summary_workbook(
    output_path: Path,
    overview_headers: list[str],
    detail_headers: list[str],
    expected_categories: list[str],
) -> None:
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        required_sheets = ["今日总览", "分类明细"]
        missing_sheets = [sheet for sheet in required_sheets if sheet not in workbook.sheetnames]
        if missing_sheets:
            raise ValueError(f"daily summary workbook missing sheets: {missing_sheets}")

        overview_ws = workbook["今日总览"]
        actual_overview_headers = list(next(overview_ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        if actual_overview_headers != overview_headers:
            raise ValueError(
                f"daily summary overview headers mismatch: {actual_overview_headers} != {overview_headers}"
            )

        detail_ws = workbook["分类明细"]
        actual_detail_headers = list(next(detail_ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        if actual_detail_headers != detail_headers:
            raise ValueError(
                f"daily summary detail headers mismatch: {actual_detail_headers} != {detail_headers}"
            )

        actual_categories = [
            row[0]
            for row in detail_ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] is not None
        ]
        if actual_categories != expected_categories:
            raise ValueError("daily summary detail category order mismatch in saved workbook")
    finally:
        workbook.close()


def save_daily_summary_workbook(
    target_date: str,
    daily_summary_df: pd.DataFrame,
    day_detail_df: pd.DataFrame,
    category_config_df: pd.DataFrame,
) -> Path:
    output_path = get_daily_summary_output_path(target_date)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    overview_df = _build_daily_overview_df(target_date, daily_summary_df)
    detail_df = _build_daily_detail_df(day_detail_df, category_config_df)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        overview_df.to_excel(writer, sheet_name="今日总览", index=False)
        detail_df.to_excel(writer, sheet_name="分类明细", index=False)

    workbook = load_workbook(output_path)
    _fit_worksheet_style(workbook["今日总览"])
    _fit_worksheet_style(workbook["分类明细"], left_headers={"类别"})
    workbook.save(output_path)
    workbook.close()

    expected_categories = category_config_df["category"].astype(str).tolist() + ["总计"]
    _verify_saved_daily_summary_workbook(
        output_path=output_path,
        overview_headers=overview_df.columns.tolist(),
        detail_headers=detail_df.columns.tolist(),
        expected_categories=expected_categories,
    )
    print(f"[daily_summary] workbook verified = ok ({output_path.name})")
    return output_path


def get_day_detail_for_date(target_date: str, detail_df: pd.DataFrame) -> pd.DataFrame:
    if detail_df.empty:
        return pd.DataFrame(columns=SCORE_DETAIL_COLUMNS)
    out = detail_df[detail_df["date"] == target_date].copy()
    if out.empty:
        return pd.DataFrame(columns=SCORE_DETAIL_COLUMNS)
    return out[SCORE_DETAIL_COLUMNS]
